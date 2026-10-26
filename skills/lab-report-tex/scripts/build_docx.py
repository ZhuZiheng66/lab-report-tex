# -*- coding: utf-8 -*-
"""Word 版报告：不需要 LaTeX。把 report.md 写进学校的 Word 模板（或新建一份），验收通过后输出 .docx（可选再转 PDF）。

和 PDF 版（build.py，LaTeX，排版最好）共用同一套数据：data/*.csv、fig/、results.tex（结果宏）、data_status.json、
requirements.md。区别只是最后一步排版用 Word。

用法：
  python build_docx.py report.md                                   # 没有模板：新建一份规范的 Word 报告
  python build_docx.py report.md --template 实验报告模版.docx         # 填进学校模板（保留模板原有栏目和样式）
  python build_docx.py report.md --template t.docx --pdf --name "39-9-张三-{expname}" --outdir ../成品
  python build_docx.py report.md --cover-later                     # 封面个人信息由用户自己填：空着不算错

report.md 写法（UTF-8）：
  ---                                  ← 开头的封面信息（可选；没给的字段在模板里保持原样）
  expname: 霍尔效应及其应用
  author: 张三
  date: 2026年9月23日
  ---
  # 实验目的                            ← 一级标题 = 模板里的章节名（不写“一、”，模板里的编号保留）
  1．了解……
  ## （一）霍尔效应                     ← 小标题
  !式 V_{H} = K_{H} I_{s} B　　（1）    ← 居中公式行；X_{下标} X^{上标} **加粗**
  !表 表1　V_{H}-I_{s} 关系 | data/t_is.csv [| 宽度cm]
  !图 fig/vh_is.pdf | 图1　V_{H}-I_{s} 关系曲线 [| 宽度cm]   ← PDF 图会自动转成 PNG（需要 PyMuPDF）
  !评分 预习/20,操作及记录/40,…         ← 只在“没有模板”时生成评分表；有模板时用模板自带的
  !照片 原始数据 | raw/a.jpg, raw/b.jpg  ← 放在文档最后，每张单独一页
  正文里的数字用 {{宏名}} 引用 results.tex 的结果（如 {{BVal}}），不要手抄。

验收（不通过就只出 build/candidate/<名字>.candidate.docx，退出码 2）：
  数据状态（data_status.json 全部冻结且未改动）、requirements.md 的章节/图/表、{{宏}} 都能解析、
  模板里能找到每个章节、没有 TODO/……/尖括号占位、没有手抄 results.tex 已有的数值、封面不空（--cover-later 可放宽）。
退出码：0 已交付；2 验收不通过；3 缺依赖（python-docx）。
依赖：python-docx；可选 PyMuPDF（PDF 图转 PNG）、Pillow、Word 或 LibreOffice（--pdf）。
"""
import argparse
import csv
import io
import os
import re
import shutil
import sys
from xml.sax.saxutils import escape

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import data_status  # noqa: E402

try:
    import docx
    from docx.enum.text import WD_BREAK
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph
except ImportError:
    print('【环境缺失】需要 python-docx：python -m pip install python-docx')
    sys.exit(3)

EXIT_OK, EXIT_CHECK, EXIT_ENV = 0, 2, 3
NS = nsdecls('w')
BODY_PPR = '<w:pPr><w:spacing w:line="300" w:lineRule="auto"/><w:ind w:firstLine="420"/></w:pPr>'
PLAIN_PPR = '<w:pPr><w:spacing w:line="300" w:lineRule="auto"/></w:pPr>'
CENTER_PPR = '<w:pPr><w:spacing w:before="60" w:after="60"/><w:jc w:val="center"/></w:pPr>'
FONT = '<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体" w:hint="eastAsia"/>'
COVER_LABELS = {
    'coursename': ['课程名称'], 'expname': ['实验名称'], 'college': ['学院', '学 院', '学    院'],
    'teacher': ['指导教师'], 'author': ['报告人', '姓名', '姓 名'], 'group': ['组号'], 'studentid': ['学号'],
    'location': ['实验地点'], 'date': ['实验时间', '实验日期'], 'submitdate': ['提交时间'], 'coursecode': ['课程编号'],
    'major': ['专业'], 'partner': ['同组者'],
}
TOKEN = re.compile(r'(\*\*.+?\*\*|_\{.*?\}|\^\{.*?\})')
NUMHEAD = re.compile(r'^[一二三四五六七八九十]+\s*[、.．]\s*')


# ---------------- 读 report.md ----------------

def parse_md(path):
    text = io.open(path, encoding='utf-8-sig').read()
    # 去掉 <!-- 注释 -->（可跨行），保留换行数，行号仍对得上；lrt:literal 标记留在原行
    text = re.sub(r'<!--.*?-->', lambda m: ('lrt:literal' if 'lrt:literal' in m.group(0) else '')
                  + '\n' * m.group(0).count('\n'), text, flags=re.S)
    meta = {}
    m = re.match(r'^---\s*\n(.*?)\n---\s*\n', text, re.S)
    if m:
        for ln in m.group(1).splitlines():
            if ':' in ln or '：' in ln:
                k, v = re.split(r'[:：]', ln, 1)
                meta[k.strip()] = v.strip()
        text = text[m.end():]
    sections, cur = [], None
    for i, ln in enumerate(text.splitlines(), 1 + (m.group(0).count('\n') if m else 0)):
        if re.match(r'^#\s+\S', ln):
            cur = {'title': ln[1:].strip(), 'lines': [], 'line': i}
            sections.append(cur)
        elif cur is not None:
            cur['lines'].append((i, ln))
        elif ln.strip() and not ln.strip().startswith('<!--'):
            sections.append({'title': '', 'lines': [(i, ln)], 'line': i})
            cur = sections[-1]
    return meta, sections


def load_macros(base, names=('results.tex',)):
    vals = {}
    for n in names:
        p = os.path.join(base, n)
        if os.path.exists(p):
            for k, v in re.findall(r'^\\def\\([A-Za-z]+)\{([^{}]*)\}', io.open(p, encoding='utf-8').read(), re.M):
                vals[k] = v
    return vals


def fmt_value(v):
    """宏值里的科学计数法转成 Word 里好看的写法：-2.647e+04 → −2.647×10^{4}。"""
    m = re.fullmatch(r'([-+]?\d*\.?\d+)[eE]([-+]?\d+)', v.strip())
    if m:
        return '%s×10^{%d}' % (m.group(1).replace('-', '−'), int(m.group(2)))
    return v.replace('-', '−') if re.fullmatch(r'-\d.*', v) else v


# ---------------- 生成 Word 元素 ----------------

LATEX_SYM = {r'\mu': 'μ', r'\Omega': 'Ω', r'\alpha': 'α', r'\beta': 'β', r'\lambda': 'λ', r'\theta': 'θ',
             r'\Delta': 'Δ', r'\rho': 'ρ', r'\omega': 'ω', r'\varphi': 'φ', r'\phi': 'φ', r'\sigma': 'σ',
             r'\pi': 'π', r'\cdot': '·', r'\times': '×', r'\pm': '±', r'\approx': '≈', r'\lg': 'lg', r'\ln': 'ln',
             r'\,': ' ', r'\%': '%', r'\sim': '~'}


def latex_lite(s):
    """把 CSV 表头/表题里常见的简单 LaTeX 转成 Word 标记（_{下标} ^{上标}）。复杂公式请在 report.md 里直接写。"""
    if '$' not in s and '\\' not in s:
        return s
    s = re.sub(r'\\(?:mathrm|text|textrm|mathit)\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\sqrt\{([^{}]*)\}', r'√(\1)', s)
    for k, v in sorted(LATEX_SYM.items(), key=lambda kv: -len(kv[0])):
        s = s.replace(k + ' ', v).replace(k, v)
    s = re.sub(r'([_^])([A-Za-z0-9])', r'\1{\2}', s)     # x_f → x_{f}
    s = s.replace('$', '')
    return re.sub(r'\\[A-Za-z]+', '', s)


def runs_xml(text, size=21, bold=False, _raw=False):
    if not _raw:
        text = latex_lite(text.replace('lrt:literal', '').rstrip())
    out = []
    for tok in TOKEN.split(text):
        if not tok:
            continue
        b, va = bold, None
        if tok.startswith('**') and tok.endswith('**'):
            out.append(runs_xml(tok[2:-2], size, True, _raw=True))   # 加粗里还可以有上下标
            continue
        elif tok.startswith('_{'):
            tok, va = tok[2:-1], 'subscript'
        elif tok.startswith('^{'):
            tok, va = tok[2:-1], 'superscript'
        rpr = FONT + ('<w:b/>' if b else '') + ('<w:vertAlign w:val="%s"/>' % va if va else '') + \
            '<w:sz w:val="%d"/><w:szCs w:val="%d"/>' % (size, size)
        out.append('<w:r><w:rPr>%s</w:rPr><w:t xml:space="preserve">%s</w:t></w:r>' % (rpr, escape(tok)))
    return ''.join(out)


def add_body(doc, el):
    """正文末尾追加——必须插在 sectPr（节属性）前面，否则内容会跑到分节符之后，页序错乱。"""
    body = doc.element.body
    sp = body.find(qn('w:sectPr'))
    if sp is not None:
        sp.addprevious(el)
    else:
        body.append(el)


def page_break():
    return parse_xml('<w:p %s><w:r><w:br w:type="page"/></w:r></w:p>' % NS)


def para(text, ppr=BODY_PPR, size=21, bold=False):
    return parse_xml('<w:p %s>%s%s</w:p>' % (NS, ppr, runs_xml(text, size, bold)))


def table_el(rows, width_cm=14.0):
    ncol = max(len(r) for r in rows)
    w = int(width_cm * 567 / ncol)
    border = ''.join('<w:%s w:val="single" w:sz="4" w:space="0" w:color="auto"/>' % s
                     for s in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'))
    xml = ['<w:tbl %s><w:tblPr><w:tblW w:w="0" w:type="auto"/><w:jc w:val="center"/>'
           '<w:tblBorders>%s</w:tblBorders><w:tblLayout w:type="fixed"/></w:tblPr><w:tblGrid>' % (NS, border)]
    xml += ['<w:gridCol w:w="%d"/>' % w] * ncol + ['</w:tblGrid>']
    for i, r in enumerate(rows):
        xml.append('<w:tr><w:trPr><w:cantSplit/>%s<w:jc w:val="center"/></w:trPr>' % ('<w:tblHeader/>' if i == 0 else ''))
        for c in range(ncol):
            cell = (r[c] if c < len(r) else '').strip()
            if re.fullmatch(r'-\d.*', cell):
                cell = '−' + cell[1:]
            xml.append('<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/><w:vAlign w:val="center"/></w:tcPr>'
                       '<w:p><w:pPr><w:jc w:val="center"/></w:pPr>%s</w:p></w:tc>' % (w, runs_xml(cell, 18, i == 0)))
        xml.append('</w:tr>')
    xml.append('</w:tbl>')
    return parse_xml(''.join(xml))


def image_source(path, workdir):
    """PDF 图（LaTeX 版用的矢量图）自动转 PNG；返回可插入 Word 的图片路径。"""
    if path.lower().endswith('.pdf'):
        png = os.path.join(workdir, 'build', 'docx-img', os.path.splitext(os.path.basename(path))[0] + '.png')
        if not os.path.exists(png) or os.path.getmtime(png) < os.path.getmtime(path):
            try:
                import fitz
            except ImportError:
                raise RuntimeError('图 %s 是 PDF，转 PNG 需要 PyMuPDF（pip install pymupdf），或让 fit_plot.py 直接输出 .png' % path)
            os.makedirs(os.path.dirname(png), exist_ok=True)
            pdf = fitz.open(path)          # 必须先拿住文档对象，否则页面会变成“孤儿”
            pdf[0].get_pixmap(dpi=220).save(png)
            pdf.close()
        return png
    return path


def image_para(doc, path, width_cm):
    p = parse_xml('<w:p %s><w:pPr><w:keepNext/><w:jc w:val="center"/></w:pPr></w:p>' % NS)
    Paragraph(p, doc._body).add_run().add_picture(path, width=Cm(width_cm))
    return p


# ---------------- 渲染章节 ----------------

def render_lines(doc, lines, base, macros, problems, extras):
    """把一个章节的标记行变成 Word 元素列表。"""
    out = []
    for lineno, raw in lines:
        s = raw.strip()
        if not s or s.startswith('<!--'):
            continue

        def sub(m):
            k = m.group(1)
            if k not in macros:
                problems.append('report.md 第 %d 行：{{%s}} 在 results.tex 里没有' % (lineno, k))
                return '{{%s}}' % k
            return fmt_value(macros[k])
        s = re.sub(r'\{\{\s*([A-Za-z]+)\s*\}\}', sub, s)
        try:
            if s.startswith('## '):
                out.append(para(s[3:].strip(), PLAIN_PPR, bold=True))
            elif s.startswith('!式'):
                out.append(para(s[2:].strip(), CENTER_PPR))
            elif s.startswith('!表'):
                parts = [x.strip() for x in s[2:].split('|')]
                p = os.path.join(base, parts[1])
                rows = [r for r in csv.reader(io.open(p, encoding='utf-8-sig')) if any(c.strip() for c in r)]
                width = float(parts[2]) if len(parts) > 2 else 14.0
                out += [para(parts[0], CENTER_PPR.replace('<w:jc', '<w:keepNext/><w:jc'), 18), table_el(rows, width)]
                extras['tables'].add(os.path.splitext(os.path.basename(parts[1]))[0])
            elif s.startswith('!图'):
                parts = [x.strip() for x in s[2:].split('|')]
                src = image_source(os.path.join(base, parts[0]), base)
                width = float(parts[2]) if len(parts) > 2 else 12.0
                out += [image_para(doc, src, width), para(parts[1] if len(parts) > 1 else '', CENTER_PPR, 18)]
                extras['figs'].add(os.path.splitext(os.path.basename(parts[0]))[0])
            elif s.startswith('!评分'):
                extras['grading'] = s[3:].strip()
            elif s.startswith('!照片'):
                title, _, files = s[3:].partition('|')
                extras['photos'].append((title.strip(), [f.strip() for f in files.split(',') if f.strip()]))
            elif s.startswith('!空'):
                out.append(para(''))
            else:
                out.append(para(s))
        except (OSError, IndexError, RuntimeError, ValueError) as e:
            problems.append('report.md 第 %d 行处理失败：%s' % (lineno, e))
    return out


def ptext(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def norm_title(t):
    return NUMHEAD.sub('', t.replace(' ', '').replace('　', '')).rstrip('：:')


def walk_cells(el):
    for child in el.iterchildren():
        if child.tag == qn('w:tbl'):
            for tr in child.iterchildren(qn('w:tr')):
                for tc in tr.iterchildren(qn('w:tc')):
                    yield tc
                    for x in walk_cells(tc):
                        yield x


def find_section(doc, title):
    """模板里的章节：先找“第一段以章节名开头”的单元格（深大这类大表格模板），再找正文里的标题段。"""
    key = norm_title(title)
    cells = []
    for tc in walk_cells(doc.element.body):
        ps = [p for p in tc.iterchildren(qn('w:p')) if ptext(p).strip()]
        if ps:
            cells.append((tc, ps[0], norm_title(ptext(ps[0]))))
    for tc, head, h in cells:
        if h.startswith(key):
            return 'cell', tc, head
    # 名称不完全一样时（老师要求写“数据记录和处理”，模板栏目叫“数据处理”）：模板栏目名被章节名包含即可
    def in_order(short, long):          # short 的每个字按顺序出现在 long 里：“数据处理” ⊂ “数据记录和处理”
        it = iter(long)
        return all(ch in it for ch in short)
    for tc, head, h in cells:
        hk = re.split(r'[（(]', h)[0]
        if 3 <= len(hk) <= 12 and in_order(hk, key) and hk[:2] == key[:2] and not re.search(r'\d+分', h):
            return 'cell', tc, head
    for p in doc.element.body.iterchildren(qn('w:p')):
        t = ptext(p).strip()
        if t and len(t) <= 30 and norm_title(t).startswith(key):
            return 'body', doc.element.body, p
    return None, None, None


def fill_section(kind, container, head, elements, all_titles):
    if kind == 'cell':
        after = False
        for el in list(container.iterchildren()):
            if el is head:
                after = True
                continue
            if after and el.tag in (qn('w:p'), qn('w:tbl')):
                container.remove(el)
        if elements and elements[-1].tag == qn('w:tbl'):
            elements.append(para(''))
        for el in elements:
            container.append(el)
    else:   # 正文标题段：删到下一个章节标题为止，再插入
        nxt = head.getnext()
        while nxt is not None and nxt.tag in (qn('w:p'), qn('w:tbl')):
            t = norm_title(ptext(nxt).strip()) if nxt.tag == qn('w:p') else ''
            if t and any(t.startswith(norm_title(x)) for x in all_titles):
                break
            following = nxt.getnext()
            container.remove(nxt)
            nxt = following
        anchor = head
        for el in elements:
            anchor.addnext(el)
            anchor = el


def replace_range(ts, lo, hi, new):
    """把段落里 [lo, hi) 这段文字（可能跨多个 run）换成 new，保留第一个 run 的格式（如下划线）。"""
    pos, first = 0, True
    for t in ts:
        s = t.text or ''
        a, b = pos, pos + len(s)
        pos = b
        if lo == hi and a <= lo <= b and first:          # 空段：插在这个位置
            t.text = s[:lo - a] + new + s[lo - a:]
            first = False
            break
        if b <= lo or a >= hi:
            continue
        a0, a1 = max(lo, a) - a, min(hi, b) - a
        t.text = s[:a0] + (new if first else '') + s[a1:]
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        first = False
    return not first


ALL_LABELS = sorted({l for ls in COVER_LABELS.values() for l in ls}, key=len, reverse=True)
BLANK = ' 　_\t'


def fill_cover(doc, meta):
    """封面：只往“标签：”后面的空白/下划线里填；模板已经写了内容的字段不动。
    日期字段是“____年____月____日”时，把年、月、日分别填进三个空。返回 (已填, 模板已有而保留)。"""
    done, kept = [], []
    paras = list(doc.element.body.iter(qn('w:p')))
    for key, labels in COVER_LABELS.items():
        val = meta.get(key)
        if not val:
            continue
        for p in paras:
            ts = list(p.iter(qn('w:t')))
            full = ''.join(t.text or '' for t in ts)
            hit = next((l for l in labels if l in full), None)
            if not hit:
                continue
            start = full.index(hit) + len(hit)
            if full[start:start + 1] in ('：', ':'):
                start += 1
            end = len(full)
            for other in ALL_LABELS:          # 同一行后面还有别的标签（“报告人：  组号：”）
                k = full.find(other, start)
                if k >= 0 and other not in labels:
                    end = min(end, k)
            seg = full[start:end]
            ymd = re.fullmatch(r'\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*', val)
            if key in ('date', 'submitdate') and '年' in seg and '月' in seg and '日' in seg:
                if not ymd:
                    break
                if re.search(r'\d', seg):
                    kept.append(key)
                    break
                spots = [m.span(1) for m in re.finditer(r'([%s]*)(年|月|日)' % BLANK, seg)][:3]
                for (lo, hi), v in reversed(list(zip(spots, ymd.groups()))):
                    width = hi - lo
                    replace_range(ts, start + lo, start + hi, v.center(max(width, len(v))))
                done.append(key)
                break
            if seg.strip(BLANK):              # 模板已经有内容（如预填的课程名、老师）
                kept.append(key)
                break
            width = len(seg.rstrip('\r\n'))
            new = val if width <= len(val) + 2 else val.center(width)
            if replace_range(ts, start, start + width, new):
                done.append(key)
            break
    return done, kept


def new_document(meta):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = sec.bottom_margin = Cm(2.4)
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    title = (meta.get('school', '') + '实验报告').strip()
    add_body(doc, parse_xml('<w:p %s><w:pPr><w:spacing w:before="1600" w:after="400"/><w:jc w:val="center"/></w:pPr>%s</w:p>'
                                      % (NS, runs_xml(title, 52, True))))
    add_body(doc, parse_xml('<w:p %s><w:pPr><w:spacing w:after="900"/><w:jc w:val="center"/></w:pPr>%s</w:p>'
                                      % (NS, runs_xml(meta.get('expname', ''), 32))))
    for key, label in (('coursename', '课程名称'), ('college', '学院'), ('author', '姓名'), ('studentid', '学号'),
                       ('group', '组号'), ('teacher', '指导教师'), ('location', '实验地点'), ('date', '实验时间'),
                       ('submitdate', '提交时间')):
        if key in meta or key in ('author', 'date'):
            add_body(doc, parse_xml(
                '<w:p %s><w:pPr><w:spacing w:line="420" w:lineRule="auto"/><w:ind w:left="2100"/></w:pPr>%s</w:p>'
                % (NS, runs_xml('%s：%s' % (label, meta.get(key, '')), 28))))
    add_body(doc, page_break())
    return doc


def grading_table(spec):
    items = [x.strip() for x in spec.split(',') if x.strip()]
    head = ['%s\n（%s分）' % tuple(i.split('/', 1)) if '/' in i else i for i in items] + ['报告整体印象', '总分']
    return [para('成绩评定', PLAIN_PPR, 24, True), table_el([[h.replace('\n', '') for h in head], [''] * len(head)], 16.0),
            para('')]


def append_photos(doc, photos, base, problems):
    for title, files in photos:
        for f in files:
            p = os.path.join(base, f)
            if not os.path.exists(p):
                problems.append('照片不存在：%s' % p)
                continue
            add_body(doc, page_break())
            if title:
                add_body(doc, para(title, CENTER_PPR, 28, True))
            add_body(doc, image_para(doc, p, 16.0))


# ---------------- 验收 ----------------

def content_checks(md_path, sections, meta, macros, required, cover_later, extras, req_figs, req_tabs):
    problems = []
    titles = [s['title'] for s in sections if s['title']]
    pos = 0
    for item in required:
        hit = next((i for i in range(pos, len(titles)) if item in titles[i]), None)
        if hit is None:
            if item in ('成绩评定', '指导教师批阅意见', '原始数据', '预习试卷') and extras.get('template'):
                continue        # 这些栏目在模板里本来就有
            if item == '指导教师批阅意见':
                continue        # 没有模板时自动生成
            if item == '原始数据' and extras['photos']:
                continue
            if item == '成绩评定' and extras.get('grading'):
                continue
            problems.append('要求的章节“%s”%s' % (item, '顺序不对' if any(item in t for t in titles) else '缺失'))
        else:
            pos = hit + 1
    for lineno, ln in ((n, l) for s in sections for n, l in s['lines']):
        body = re.sub(r'\{\{[^}]*\}\}', '', ln)
        if re.search(r'……|TODO|待补|待填|XXX|<[^<>\n]{1,12}>', body):
            problems.append('report.md 第 %d 行还有占位符/TODO' % lineno)
        if 'lrt:literal' in ln or ln.strip().startswith(('!表', '!图', '!照片', '!评分')):
            continue
        for numtxt in re.findall(r'(?<![\d.])(\d*\.\d+)(?![\d.])', body):
            for k, v in macros.items():
                vv = v.strip().lstrip('+-')
                m = re.fullmatch(r'(\d*\.\d+)(?:[eE][-+]?\d+)?', vv)
                if m and m.group(1) == numtxt and len(re.sub(r'^[0.]+', '', numtxt.replace('.', ''))) >= 3:
                    problems.append('report.md 第 %d 行手写了 %s，与宏 {{%s}} 相同：请改用 {{%s}}' % (lineno, numtxt, k, k))
                    break
    for f in req_figs:
        if os.path.splitext(os.path.basename(f))[0] not in extras['figs']:
            problems.append('需求清单要求的图没有出现在 report.md：%s' % f)
    for t in req_tabs:
        if os.path.splitext(os.path.basename(t))[0] not in extras['tables']:
            problems.append('需求清单要求的表没有出现在 report.md：%s（!表 引用同名 CSV）' % t)
    empty = [k for k in ('expname', 'author', 'date') if not meta.get(k)]
    if empty and not cover_later and not extras.get('template'):
        problems.append('封面信息为空：%s（不知道就问用户；用户要自己填就加 --cover-later）' % ', '.join(empty))
    return problems


def to_pdf(docx_path):
    import convert_legacy
    pdf = os.path.splitext(docx_path)[0] + '.pdf'
    try:
        convert_legacy.office_convert(os.path.abspath(docx_path), os.path.abspath(pdf), 'word')
    except Exception:
        convert_legacy.lo_convert(os.path.abspath(docx_path), os.path.abspath(pdf))
    return pdf


def main():
    import build   # 复用需求清单解析
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('md', nargs='?', default='report.md')
    ap.add_argument('--template', help='学校的 Word 模板（.docx；.doc 先用 convert_legacy.py 转换）')
    ap.add_argument('--require', help='章节顺序，逗号分隔；不给则读 requirements.md')
    ap.add_argument('--name', help='输出文件名（不含扩展名），可用 {expname} 等封面字段')
    ap.add_argument('--outdir', default='.')
    ap.add_argument('--pdf', action='store_true', help='再用 Word / LibreOffice 转一份 PDF')
    ap.add_argument('--cover-later', action='store_true', help='封面个人信息由用户自己填，空着不算错')
    ap.add_argument('--allow-unfrozen', action='store_true', help='原始数据未冻结时只警告（仅草稿）')
    a = ap.parse_args()

    md = os.path.abspath(a.md)
    base = os.path.dirname(md)
    meta, sections = parse_md(md)
    macros = load_macros(base)
    tpl = a.template or meta.get('template')
    if tpl and not os.path.isabs(tpl):
        tpl = os.path.join(base, tpl)
    if tpl and tpl.lower().endswith('.doc'):
        print('✘ 模板是 .doc 老格式，先运行：python convert_legacy.py "%s"' % tpl)
        sys.exit(EXIT_CHECK)

    problems, warnings = [], []
    ds = data_status.check(base)
    (warnings if a.allow_unfrozen else problems).extend('数据状态：' + p for p in ds)
    req_sections, req_figs, req_tabs = build.parse_requirements(os.path.join(base, 'requirements.md'))
    required = [s.strip() for s in re.split(r'[,，]', a.require or req_sections or '') if s.strip()]

    doc = docx.Document(tpl) if tpl else new_document(meta)
    extras = {'figs': set(), 'tables': set(), 'photos': [], 'grading': None, 'template': bool(tpl)}
    all_titles = [s['title'] for s in sections if s['title']]
    for s in sections:
        els = render_lines(doc, s['lines'], base, macros, problems, extras)
        if not s['title']:
            if els:
                problems.append('report.md 第 %d 行的内容不在任何 # 章节下面' % s['line'])
            continue
        if tpl:
            kind, container, head = find_section(doc, s['title'])
            if kind is None:
                problems.append('模板里找不到章节“%s”（report.md 第 %d 行）；先用 read_template.py 看模板的章节名'
                                % (s['title'], s['line']))
                continue
            tname = ptext(head).strip()
            if norm_title(tname) != norm_title(s['title']):
                print('  「%s」填进了模板的「%s」栏（模板栏目名不改）' % (s['title'], tname))
            fill_section(kind, container, head, els, all_titles)
        else:
            idx = sum(1 for x in all_titles[:all_titles.index(s['title']) + 1])
            heading = '%s、%s' % ('一二三四五六七八九十'[idx - 1] if idx <= 10 else str(idx), s['title'])
            add_body(doc, para(heading, PLAIN_PPR, 24, True))
            for el in els:
                add_body(doc, el)
    if tpl:
        filled, kept = fill_cover(doc, meta)
        if filled:
            print('封面已填：%s' % '、'.join(filled))
        if kept:
            print('封面这些栏模板里已有内容，保持不动：%s' % '、'.join(kept))
    else:
        add_body(doc, para('指导教师批阅意见', PLAIN_PPR, 24, True))
        for _ in range(5):
            add_body(doc, para(''))
        if extras['grading']:
            for el in grading_table(extras['grading']):
                add_body(doc, el)
    append_photos(doc, extras['photos'], base, problems)
    # 文档必须以 sectPr 结尾：把它移回最后
    body = doc.element.body
    sp = body.find(qn('w:sectPr'))
    if sp is not None:
        body.remove(sp)
        body.append(sp)

    problems += content_checks(md, sections, meta, macros, required, a.cover_later, extras, req_figs, req_tabs)

    keys = dict(meta)
    if a.name:
        try:
            name = a.name.format(**keys)
        except KeyError as e:
            raise SystemExit('--name 里的 %s 在 report.md 开头的封面信息里没有' % e)
    else:
        name = meta.get('expname') or os.path.splitext(os.path.basename(md))[0]
    name = re.sub(r'[\\/:*?"<>|]', '_', name)

    if problems:
        cand_dir = os.path.join(base, 'build', 'candidate')
        os.makedirs(cand_dir, exist_ok=True)
        cand = os.path.join(cand_dir, name + '.candidate.docx')
        doc.save(cand)
        print('✘ 验收未通过（%d 项），没有输出正式文件：' % len(problems))
        for p in problems:
            print('   -', p)
        for w in warnings:
            print('   ! 警告：', w)
        print('候选稿（仅供检查，不是交付结果）：%s' % cand)
        sys.exit(EXIT_CHECK)

    os.makedirs(a.outdir, exist_ok=True)
    final = os.path.abspath(os.path.join(a.outdir, name + '.docx'))
    doc.save(final)
    for w in warnings:
        print('! 警告：', w)
    if not warnings:
        data_status.mark_used(base)
    print('✔ 验收通过，输出：%s（%.0f KB）' % (final, os.path.getsize(final) / 1024))
    if a.pdf:
        try:
            print('✔ PDF：%s' % to_pdf(final))
        except Exception as e:
            print('! 没能转 PDF（%s）：docx 已生成，可以用 Word/WPS 打开后另存为 PDF' % e)
    print('提示：Word 版公式是普通文字加上下标；需要更好的公式和排版，用 build.py（LaTeX/PDF 版）。')
    sys.exit(EXIT_OK)


if __name__ == '__main__':
    main()
