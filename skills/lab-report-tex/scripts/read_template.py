# -*- coding: utf-8 -*-
"""解析学校的 Word 报告模板（.docx）：封面字段、章节顺序、评分栏、格式注释，**包括嵌套表格**。

为什么需要它：学校模板常把整份报告做成一张大表，“成绩评定”里再嵌一张评分表。只遍历
document.tables 时，嵌套表完全看不到——评分栏就会被漏掉或自创。本脚本递归遍历每个单元格里的
段落和嵌套表格（cell.tables），按文档顺序输出全部内容。

用法：
  python read_template.py 模板.docx                 # 打印结构 + 建议的 --require 和 \\labgrading
  python read_template.py 模板.docx --dump          # 另外按顺序打印全部文字（含嵌套表，标出层级）
.doc 先转换：python convert_legacy.py 模板.doc
依赖：python-docx（pip install python-docx）。
"""
import argparse
import re
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

try:
    import docx
    from docx.oxml.ns import qn
except ImportError:
    raise SystemExit('需要 python-docx：pip install python-docx')

SECTION_WORDS = ['实验目的', '实验原理', '实验仪器', '实验内容', '实验步骤', '数据记录', '数据处理', '实验数据',
                 '结果陈述', '实验结果', '结果与讨论', '实验结论', '实验总结', '讨论', '思考题', '误差分析',
                 '指导教师批阅意见', '教师评阅', '批阅意见', '成绩评定', '原始数据']
COVER_RE = re.compile(r'(课程名称|课程编号|实验名称|学\s*院|专\s*业|班\s*级|指导教师|报\s*告\s*人|姓\s*名|学\s*号|'
                      r'组\s*号|同组者|实验地点|实验时间|提交时间|实验日期)\s*[：:]?')
NUM_HEAD = re.compile(r'^\s*([一二三四五六七八九十]+)\s*[、.．]\s*(\S.{0,20})')


def ptext(p_el):
    return ''.join(t.text or '' for t in p_el.iter(qn('w:t')))


def is_numbered(p_el):
    """Word 自动编号的段落（显示为“二、”但文字里没有）。"""
    return p_el.find('.//' + qn('w:numPr')) is not None


def walk(el, depth=0, out=None):
    """按文档顺序产出 (类型, 深度, 内容)。类型：p 段落；row 表格行（单元格文字列表）；tbl_start/tbl_end。"""
    if out is None:
        out = []
    for child in el.iterchildren():
        tag = child.tag
        if tag == qn('w:p'):
            t = ptext(child).strip()
            if t:
                out.append(('p', depth, t, is_numbered(child)))
        elif tag == qn('w:tbl'):
            out.append(('tbl_start', depth, '', False))
            for tr in child.iterchildren(qn('w:tr')):
                cells = []
                for tc in tr.iterchildren(qn('w:tc')):
                    cells.append(' / '.join(ptext(p).strip() for p in tc.iterchildren(qn('w:p')) if ptext(p).strip()))
                out.append(('row', depth, cells, False))
                for tc in tr.iterchildren(qn('w:tc')):
                    walk(tc, depth + 1, out)       # 单元格里的段落和嵌套表格
            out.append(('tbl_end', depth, '', False))
        elif tag in (qn('w:sdt'), qn('w:sdtContent')):
            walk(child, depth, out)
    return out


def dedupe(seq):
    seen, res = set(), []
    for s in seq:
        if s not in seen:
            seen.add(s)
            res.append(s)
    return res


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('docx')
    ap.add_argument('--dump', action='store_true')
    a = ap.parse_args()
    if a.docx.lower().endswith('.doc'):
        raise SystemExit('这是 .doc 老格式，先运行：python convert_legacy.py "%s"' % a.docx)

    items = walk(docx.Document(a.docx).element.body)
    paras = [(d, t, num) for kind, d, t, num in items if kind == 'p']

    # 封面字段
    cover = dedupe(m.group(1).replace(' ', '').replace('　', '')
                   for _, t, _ in paras for m in COVER_RE.finditer(t))

    # 章节：显式编号“一、xx”，或自动编号段落 / 单元格开头是已知章节名
    sections = []
    after_grade = False
    for d, t, num in paras:
        if re.search(r'\d+\s*分', t):          # 评分栏的格子（“数据处理与结果陈述（30分）”）不是章节
            continue
        if after_grade and '原始数据' not in t:  # 成绩评定之后只剩附页
            continue
        if t.startswith('成绩评定'):
            after_grade = True
        m = NUM_HEAD.match(t)
        name = None
        if m:
            name = m.group(2)
        elif num or len(t) <= 14:
            for w in SECTION_WORDS:
                if t.startswith(w):
                    name = t
                    break
        if name:
            name = re.split(r'[：:（(]', name)[0].strip()
            if name and not any(name == s for s in sections):
                sections.append(name)

    # 评分栏：找“成绩评定”之后最近的一张（可能是嵌套的）表的第一行
    grading_header, nested = None, False
    seen_grade = False
    for i, (kind, d, t, _) in enumerate(items):
        if kind == 'p' and '成绩评定' in t:
            seen_grade = True
        elif kind == 'row' and any('成绩评定' in c for c in t):
            seen_grade = True
            # 同一行里若只有一个单元格，评分表一定在它的嵌套表里，继续往后找
            if len(t) > 1 and any(re.search(r'\d+\s*分', c) for c in t):
                grading_header, nested = t, d > 0
                break
        elif seen_grade and kind == 'row' and len(t) >= 3:
            grading_header, nested = t, d > 0
            break

    notes = [t for _, t, _ in paras if re.match(r'^\s*(注|说明|备注)\s*[：:]', t) or '单独起页' in t or '号字' in t]

    print('== 封面字段 ==')
    print('  ' + '、'.join(cover) if cover else '  （没识别到，可能是图片或文本框，请看原文件）')
    print('\n== 章节顺序 ==')
    for i, s in enumerate(sections, 1):
        print('  %d. %s' % (i, s))
    print('\n== 评分栏（逐字照抄，不要自创、不要省略）==')
    if grading_header:
        cells = [re.sub(r'\s*/\s*', '', c).strip() for c in grading_header]
        print('  %s%s' % (' | '.join(cells), '   ← 嵌套表格' if nested else ''))
        scored, extra = [], []
        for c in cells:
            m = re.match(r'^(.*?)[（(]?\s*(\d+)\s*分[）)]?$', c.replace(' ', ''))
            if m:
                scored.append('%s/%s' % (m.group(1).strip('（( '), m.group(2)))
            elif c:
                extra.append(c.replace(' ', ''))
        print('\n  建议：\\labgrading[%s]{%s}' % (','.join(extra), ','.join(scored)))
        total = sum(int(s.split('/')[1]) for s in scored)
        print('  （有分值的栏目合计 %d 分%s）' % (total, '' if total == 100 else '，不是 100，请核对原模板'))
    else:
        print('  ✘ 没找到评分表。打开原文件看“成绩评定”处——可能是图片，或格式特殊；不要自己编栏目')
    print('\n== 格式注释 / 要求 ==')
    for n in notes or ['（无）']:
        print('  ' + n)

    req = [('原始数据' if '原始数据' in s else s) for s in sections]
    print('\n建议：build.py --require "%s"' % ','.join(req))

    if a.dump:
        print('\n== 全部内容（按顺序，缩进=嵌套层级）==')
        for kind, d, t, num in items:
            ind = '  ' * (d + 1)
            if kind == 'p':
                print('%s%s%s' % (ind, '[自动编号] ' if num else '', t))
            elif kind == 'row':
                print('%s| %s |' % (ind, ' | '.join(t)))
            elif kind == 'tbl_start':
                print('%s┌ 表格%s' % (ind, '（嵌套）' if d else ''))


if __name__ == '__main__':
    main()
