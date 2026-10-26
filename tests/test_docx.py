# -*- coding: utf-8 -*-
"""Word 版（build_docx.py）：不需要 LaTeX；验收不过不出正式 .docx；能填学校模板而不改模板栏目。"""
import io
import os
import shutil
import tempfile
import unittest

from helpers import SCRIPTS, has_module, run

BUILD = os.path.join(SCRIPTS, 'build_docx.py')
DS = os.path.join(SCRIPTS, 'data_status.py')

MD = '''---
expname: 测试实验
author: 张三
date: 2026年10月1日
---
# 实验目的
测量某个量。
# 数据记录和处理
结果 k={{KVal}}，表如下：
!表 表1　数据 | data.csv
'''


def text_of(path):
    import docx
    d = docx.Document(path)
    return '\n'.join(''.join(t.text or '' for t in p.iter(docx.oxml.ns.qn('w:t')))
                     for p in d.element.body.iter(docx.oxml.ns.qn('w:p')))


@unittest.skipUnless(has_module('docx'), '没有 python-docx')
class DocxBuild(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.w('raw.csv', 'x\n1\n')
        self.w('data.csv', 'I_{s}/mA,V/mV\n1.00,-0.40\n')
        self.w('results.tex', '\\def\\KVal{1.234}%  test\n')
        run([DS, 'add', 'raw.csv'], cwd=self.tmp)
        run([DS, 'review', 'raw.csv', '--confirmed-by', '测试：模拟用户确认'], cwd=self.tmp)
        run([DS, 'freeze', 'raw.csv'], cwd=self.tmp)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def w(self, name, s):
        with io.open(os.path.join(self.tmp, name), 'w', encoding='utf-8') as f:
            f.write(s)

    def build(self, md, *extra):
        self.w('report.md', md)
        code, out = run([BUILD, 'report.md', '--outdir', 'out', '--require', '实验目的,数据记录和处理'] + list(extra),
                        cwd=self.tmp)
        o = os.path.join(self.tmp, 'out')
        files = [f for f in os.listdir(o) if f.endswith('.docx')] if os.path.isdir(o) else []
        return code, out, files

    def test_pass_without_latex(self):
        code, out, files = self.build(MD, )
        self.assertEqual(code, 0, out)
        self.assertEqual(files, ['测试实验.docx'])
        t = text_of(os.path.join(self.tmp, 'out', files[0]))
        self.assertIn('1.234', t, '{{宏}} 没有被替换')
        self.assertIn('−0.40', t, '负数应换成减号')
        self.assertIn('实验目的', t)

    def test_unknown_macro_blocks(self):
        code, out, files = self.build(MD.replace('{{KVal}}', '{{Nope}}'))
        self.assertEqual(code, 2, out)
        self.assertEqual(files, [])
        self.assertIn('{{Nope}}', out)
        self.assertTrue(os.listdir(os.path.join(self.tmp, 'build', 'candidate')))

    def test_hand_copied_and_todo_block(self):
        code, out, files = self.build(MD.replace('{{KVal}}', '1.234') + 'TODO\n')
        self.assertEqual(code, 2, out)
        self.assertEqual(files, [])
        self.assertIn('手写了 1.234', out)
        self.assertIn('TODO', out)

    def test_unfrozen_data_blocks(self):
        self.w('raw.csv', 'x\n2\n')
        code, out, files = self.build(MD)
        self.assertEqual(code, 2, out)
        self.assertEqual(files, [])

    def test_fill_school_template(self):
        """模拟学校模板：大表格，每章一个单元格；封面有下划线空位和预填内容。"""
        import docx
        d = docx.Document()
        d.add_paragraph('课程名称：      大学物理实验      ')
        d.add_paragraph('实验名称：                  ')
        d.add_paragraph('实验时间：      年    月    日')
        tbl = d.add_table(rows=3, cols=1)
        tbl.cell(0, 0).text = '一、实验目的'
        tbl.cell(1, 0).text = '五、数据处理'          # 栏目名比要求的“数据记录和处理”短
        tbl.cell(2, 0).text = '成绩评定'
        d.save(os.path.join(self.tmp, 'tpl.docx'))
        code, out, files = self.build(MD + '\n# 结果陈述\n见上。\n', '--template', 'tpl.docx')
        self.assertNotIn('结果陈述', ''.join(files))
        self.assertEqual(code, 2, out)                 # 模板里没有“结果陈述”：必须报出来，不能悄悄丢掉
        self.assertIn('模板里找不到章节“结果陈述”', out)
        code, out, files = self.build(MD, '--template', 'tpl.docx')
        self.assertEqual(code, 0, out)
        t = text_of(os.path.join(self.tmp, 'out', files[0]))
        self.assertIn('五、数据处理', t, '模板栏目名不能被改')
        self.assertIn('测量某个量', t)
        self.assertIn('测试实验', t)
        self.assertEqual(t.count('大学物理实验'), 1, '模板预填的课程名不能被重复填')
        self.assertIn('2026', t)
        self.assertIn('成绩评定', t)


if __name__ == '__main__':
    unittest.main()
